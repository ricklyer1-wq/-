#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as error:
    raise SystemExit("openpyxl is required to import .xlsx card data.") from error


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "data" / "source" / "cards.xlsx"
DEFAULT_OUTPUT = ROOT / "data" / "source" / "cards.tsv"

CARD_FIELDS = [
    "id",
    "name",
    "cost",
    "type",
    "element",
    "faction",
    "pool",
    "npc_only",
    "description",
    "effect_logic",
    "rarity",
]

TYPE_MAP = {
    "攻擊": "attack",
    "技能": "skill",
    "能力": "power",
    "狀態": "status",
}

RARITY_MAP = {
    "基礎": "BASIC",
    "普通": "N",
    "罕見": "R",
    "稀有": "SR",
    "史詩": "SSR",
    "傳說": "SSR",
}

PREFIX_META = {
    "BAS": ("neutral", "base", "base"),
    "GLD": ("metal", "cangfenggu", "metal"),
    "WOD": ("wood", "kumuzhai", "wood"),
    "WAT": ("water", "xuanbinggong", "water"),
    "FIR": ("fire", "fentiangu", "fire"),
    "ERT": ("earth", "zhenyuezong", "earth"),
    "SOU": ("soul", "taixuguan", "soul"),
    "WPN": ("neutral", "qianjimen", "weapon"),
    "EVL": ("evil", "evil", "evil"),
    "STS": ("neutral", "base", "status"),
    "TKN": ("neutral", "base", "token"),
}

TOKEN_META = {
    "TKN_001": ("metal", "cangfenggu", "token"),
    "TKN_002": ("neutral", "qianjimen", "token"),
    "WOD_T01": ("wood", "kumuzhai", "token"),
}


class ImportErrorWithContext(Exception):
    pass


def main() -> int:
    parser = argparse.ArgumentParser(description="Import card data from an .xlsx workbook into cards.tsv.")
    parser.add_argument("input", nargs="?", default=str(DEFAULT_INPUT), help="Source .xlsx path.")
    parser.add_argument("--out", default=str(DEFAULT_OUTPUT), help="Output cards.tsv path.")
    args = parser.parse_args()

    source_path = Path(args.input)
    output_path = Path(args.out)
    if not source_path.is_absolute():
        source_path = ROOT / source_path
    if not output_path.is_absolute():
        output_path = ROOT / output_path

    try:
        cards = import_cards(source_path)
        write_cards_tsv(output_path, cards)
    except ImportErrorWithContext as error:
        print(f"Import error: {error}", file=sys.stderr)
        return 1

    print(f"Imported {len(cards)} cards into {output_path.relative_to(ROOT)}")
    return 0


def import_cards(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise ImportErrorWithContext(f"Missing source workbook: {path}")

    workbook = load_workbook(path, data_only=True)
    cards: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    for sheet in workbook.worksheets:
        headers = [clean_cell(cell.value) for cell in sheet[1]]
        if all(field in headers for field in CARD_FIELDS):
            sheet_cards = import_standard_sheet(sheet, headers)
        elif {"Card_ID", "名稱 (Name)", "類型", "耗能", "文本描述", "程式邏輯參考 (Effect_Logic)"}.issubset(set(headers)):
            sheet_cards = import_legacy_sheet(sheet)
        else:
            continue

        for card in sheet_cards:
            card_id = card["id"]
            if card_id in seen_ids:
                raise ImportErrorWithContext(f"Duplicate card id: {card_id}")
            seen_ids.add(card_id)
            cards.append(card)

    if not cards:
        raise ImportErrorWithContext("No card rows were found.")
    return cards


def import_standard_sheet(sheet: Any, headers: list[str]) -> list[dict[str, Any]]:
    index = {header: position for position, header in enumerate(headers)}
    cards: list[dict[str, Any]] = []

    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = {field: clean_cell(row[index[field]] if index[field] < len(row) else "") for field in CARD_FIELDS}
        card_id = raw["id"]
        if not card_id:
            continue
        card = {
            "id": card_id,
            "name": require_value(raw["name"], "name", card_id),
            "cost": normalize_cost(raw["cost"], card_id),
            "type": normalize_type(raw["type"], card_id),
            "element": require_value(raw["element"], "element", card_id),
            "faction": require_value(raw["faction"], "faction", card_id),
            "pool": require_value(raw["pool"], "pool", card_id),
            "npc_only": normalize_bool_text(raw["npc_only"], card_id),
            "description": require_value(raw["description"], "description", card_id),
            "effect_logic": require_value(raw["effect_logic"], "effect_logic", card_id),
            "rarity": normalize_rarity(card_id, raw["rarity"]),
        }
        cards.append(card)

    return cards


def import_legacy_sheet(sheet: Any) -> list[dict[str, Any]]:
    cards: list[dict[str, Any]] = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        card_id = clean_cell(row[0] if len(row) > 0 else None)
        if not card_id:
            continue

        name = clean_cell(row[1] if len(row) > 1 else None)
        type_or_cost = clean_cell(row[2] if len(row) > 2 else None)
        cost_or_type = clean_cell(row[3] if len(row) > 3 else None)
        rarity_source = clean_cell(row[4] if len(row) > 4 else None)
        description = clean_cell(row[5] if len(row) > 5 else None)
        effect_logic = clean_cell(row[6] if len(row) > 6 else None)

        if is_cost(type_or_cost) and not is_cost(cost_or_type):
            cost = normalize_cost(type_or_cost, card_id)
            card_type = normalize_type(cost_or_type, card_id)
        else:
            card_type = normalize_type(type_or_cost, card_id)
            cost = normalize_cost(cost_or_type, card_id)

        element, faction, pool = card_meta(card_id)
        cards.append(
            {
                "id": card_id,
                "name": require_value(name, "name", card_id),
                "cost": cost,
                "type": card_type,
                "element": element,
                "faction": faction,
                "pool": pool,
                "npc_only": str(card_id.startswith("EVL_")).lower(),
                "description": require_value(description, "description", card_id),
                "effect_logic": require_value(effect_logic, "effect_logic", card_id),
                "rarity": normalize_rarity(card_id, rarity_source),
            }
        )
    return cards


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_cost(value: str) -> bool:
    text = value.strip().upper()
    return text == "X" or text.isdigit()


def normalize_cost(value: str, card_id: str) -> str:
    text = value.strip().upper()
    if text == "X" or text.isdigit():
        return text
    raise ImportErrorWithContext(f"Card {card_id} has invalid cost: {value!r}")


def normalize_type(value: str, card_id: str) -> str:
    text = value.strip()
    if text in TYPE_MAP:
        return TYPE_MAP[text]
    if text in TYPE_MAP.values():
        return text
    raise ImportErrorWithContext(f"Card {card_id} has invalid type: {value!r}")


def normalize_bool_text(value: str, card_id: str) -> str:
    text = value.strip().lower()
    if text in {"true", "false", "1", "0", "yes", "no", "y", "n", "是", "否"}:
        return text
    raise ImportErrorWithContext(f"Card {card_id} has invalid npc_only value: {value!r}")


def normalize_rarity(card_id: str, value: str) -> str:
    if card_id.startswith("STS_"):
        return "STATUS"
    if card_id.startswith("TKN_") or "_T" in card_id:
        return "TOKEN"
    text = value.strip()
    return RARITY_MAP.get(text, text or "N")


def card_meta(card_id: str) -> tuple[str, str, str]:
    if card_id in TOKEN_META:
        return TOKEN_META[card_id]
    prefix = card_id.split("_", 1)[0]
    if prefix in PREFIX_META:
        return PREFIX_META[prefix]
    return ("neutral", "base", "base")


def require_value(value: str, field: str, card_id: str) -> str:
    if not value:
        raise ImportErrorWithContext(f"Card {card_id} is missing {field}.")
    return value


def write_cards_tsv(path: Path, cards: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CARD_FIELDS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(cards)


if __name__ == "__main__":
    raise SystemExit(main())
